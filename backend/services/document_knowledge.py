"""Document-as-knowledge (Phase 3 module K).

Audit finding (documented here since it drove this design): ZIP3 already
had an `ATTACHABLE_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}` constant
in api/uploads.py and a matching "attachable" file-kind classification, but
no code anywhere actually extracted text from those files, and the only
consumer of that classification (the "attach" upload action) runs the
audio/video transcription pipeline against them, which cannot work on a
PDF. So: text extraction from documents was genuinely missing, not
duplicated. Chunking, embedding, and Qdrant indexing were NOT missing —
`semantic_search_service.index_session()` already takes an arbitrary
(id, title, text, summary, org_id) tuple and is not meeting-specific in any
way that would need changing, so a document is indexed through the exact
same call meetings and lectures use, keyed by a synthetic `doc_id` instead
of a RecordingSession's session_id. No second vector store, no second
pipeline, per Phase 3B's explicit instruction.
"""
from __future__ import annotations

import hashlib
import io
import logging
import uuid
from typing import Optional

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MAX_DOCUMENT_BYTES = 20 * 1024 * 1024  # 20MB — generous for text documents, bounded to avoid unbounded memory use


class UnsupportedDocumentType(ValueError):
    pass


class DocumentTooLarge(ValueError):
    pass


def new_doc_id() -> str:
    return f"doc_{uuid.uuid4().hex}"


def content_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from a supported document. Raises
    UnsupportedDocumentType / DocumentTooLarge rather than silently
    returning empty text, so a caller can't mistake 'extraction wasn't
    attempted' for 'the document really was empty'."""
    if len(data) > MAX_DOCUMENT_BYTES:
        raise DocumentTooLarge(f"Document exceeds {MAX_DOCUMENT_BYTES} bytes")

    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocumentType(f"Unsupported document type: {ext or '(none)'}")

    if ext in (".txt", ".md"):
        return data.decode("utf-8", errors="replace")

    if ext == ".docx":
        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        parts.append(cell.text)
        return "\n".join(parts)

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
        return "\n\n".join(parts)

    raise UnsupportedDocumentType(f"Unsupported document type: {ext}")  # pragma: no cover — unreachable given the check above


def index_document(doc_id: str, title: str, text: str, organization_id: int, created_at: Optional[str] = None) -> None:
    """Index a document into the SAME Qdrant collection/service meetings
    and lectures use. `summary=None` — documents don't get an LLM summary
    at index time (nothing in Phase 3B asked for one); the full extracted
    text is what's chunked and made retrievable."""
    from services.semantic_search_service import semantic_search

    semantic_search.index_session(
        session_id=doc_id,
        title=title,
        transcript=text,
        summary=None,
        created_at=created_at,
        organization_id=organization_id,
        source_kind="document",
    )


def delete_document_index(doc_id: str) -> None:
    from services.semantic_search_service import semantic_search

    semantic_search.delete_session(doc_id)
